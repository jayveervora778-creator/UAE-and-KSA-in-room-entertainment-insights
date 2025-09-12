import pandas as pd
import openpyxl
from openpyxl import load_workbook

def investigate_excel_structure():
    """Investigate the raw Excel structure to understand data organization"""
    
    print("=== INVESTIGATING RAW EXCEL STRUCTURE ===\n")
    
    # Load workbook
    wb = load_workbook('data/survey_data.xlsx', data_only=True)
    ws = wb.active
    
    print(f"Sheet name: {ws.title}")
    print(f"Max row: {ws.max_row}")
    print(f"Max column: {ws.max_column}")
    print()
    
    # Examine first few rows to understand structure
    print("=== FIRST 10 ROWS (Raw Cell Values) ===")
    for row_idx in range(1, min(11, ws.max_row + 1)):
        row_data = []
        for col_idx in range(1, min(21, ws.max_column + 1)):  # First 20 columns
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            row_data.append(str(cell_value) if cell_value is not None else "")
        print(f"Row {row_idx}: {row_data}")
    print()
    
    # Focus on B1-A columns specifically
    print("=== LOOKING FOR B1-A PATTERN COLUMNS ===")
    
    # Read with pandas to get column names
    df = pd.read_excel('data/survey_data.xlsx')
    b1_a_columns = [col for col in df.columns if 'B1-A' in str(col)]
    
    print(f"Found B1-A columns: {b1_a_columns}")
    print()
    
    # Examine each B1-A column
    for col in b1_a_columns:
        print(f"=== COLUMN: {col} ===")
        values = df[col].dropna()
        print(f"Total responses: {len(values)}")
        print(f"Unique values: {values.nunique()}")
        print(f"Value counts:")
        print(values.value_counts().head(10))
        print(f"Sample values: {values.head(10).tolist()}")
        print()
    
    # Also examine the raw structure without pandas processing
    print("=== RAW STRUCTURE ANALYSIS ===")
    
    # Find columns that might contain B1-A data
    header_rows = []
    for row_idx in range(1, 6):  # Check first 5 rows for headers
        row_data = []
        for col_idx in range(1, ws.max_column + 1):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            row_data.append(str(cell_value) if cell_value is not None else "")
        header_rows.append(row_data)
        print(f"Header row {row_idx}: {row_data[:30]}...")  # First 30 columns
    
    print()
    
    # Look for patterns in column headers
    for row_idx, row in enumerate(header_rows, 1):
        b1_positions = []
        for col_idx, cell in enumerate(row):
            if 'B1' in str(cell) or 'A/' in str(cell):
                b1_positions.append((col_idx, cell))
        if b1_positions:
            print(f"Row {row_idx} B1-A patterns: {b1_positions}")
    
    print("\n=== ANALYSIS COMPLETE ===")

if __name__ == "__main__":
    investigate_excel_structure()
