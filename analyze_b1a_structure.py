import pandas as pd
from openpyxl import load_workbook

def analyze_b1a_responses():
    """Analyze the B1-A multi-response question structure"""
    
    print("=== ANALYZING B1-A MULTI-RESPONSE STRUCTURE ===\n")
    
    # Load workbook to examine raw structure
    wb = load_workbook('data/survey_data.xlsx', data_only=True)
    ws = wb.active
    
    # Get the legend for B1-A
    legends_row = []
    for col_idx in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=3, column=col_idx).value  # Row 3 has legends
        legends_row.append(str(cell_value) if cell_value is not None else "")
    
    headers_row = []
    for col_idx in range(1, ws.max_column + 1):
        cell_value = ws.cell(row=4, column=col_idx).value  # Row 4 has column headers
        headers_row.append(str(cell_value) if cell_value is not None else "")
    
    print("Column headers (row 4):")
    for i, header in enumerate(headers_row[:20]):
        if 'B1-A' in header:
            print(f"  Column {i+1}: {header}")
    
    # Find B1-A legend
    b1a_legend_col = None
    for i, legend in enumerate(legends_row):
        if 'Location' in legend and 'Price' in legend:
            b1a_legend_col = i
            print(f"\nB1-A Legend (Column {i+1}):")
            print(f"  {legend}")
            break
    
    # Find B1-A columns
    b1a_columns = []
    for i, header in enumerate(headers_row):
        if 'B1-A/' in header:
            b1a_columns.append((i, header))
    
    print(f"\nB1-A Response Columns:")
    for col_idx, header in b1a_columns:
        print(f"  Column {col_idx+1}: {header}")
    
    # Analyze actual responses
    print(f"\n=== ANALYZING ACTUAL RESPONSES ===")
    
    # Get a few sample rows
    sample_responses = []
    for row_idx in range(5, min(15, ws.max_row + 1)):  # First 10 response rows
        response = []
        for col_idx, header in b1a_columns:
            cell_value = ws.cell(row=row_idx, column=col_idx + 1).value
            response.append(str(cell_value) if cell_value is not None else "")
        sample_responses.append(response)
        
        print(f"Response {row_idx-4}: {response}")
    
    # Count patterns
    print(f"\n=== RESPONSE PATTERNS ===")
    
    # Load with pandas to get all data
    df = pd.read_excel('data/survey_data.xlsx', header=3)  # Use row 4 as headers
    
    b1a_cols = [col for col in df.columns if 'B1-A/' in str(col)]
    print(f"B1-A columns found: {b1a_cols}")
    
    # Analyze each column
    for col in b1a_cols:
        values = df[col].dropna()
        print(f"\nColumn {col}:")
        print(f"  Total responses: {len(values)}")
        print(f"  Unique values: {values.nunique()}")
        print(f"  Value counts:")
        print(values.value_counts().head())
    
    # Now analyze how many choices each respondent made
    print(f"\n=== MULTI-CHOICE ANALYSIS ===")
    
    # Count responses per row
    b1a_df = df[b1a_cols].copy()
    
    # Convert to numeric, treating non-numeric as NaN
    for col in b1a_cols:
        b1a_df[col] = pd.to_numeric(b1a_df[col], errors='coerce')
    
    # Count non-null responses per row
    choices_per_respondent = b1a_df.notna().sum(axis=1)
    
    print(f"Choices per respondent:")
    print(choices_per_respondent.value_counts().sort_index())
    
    print(f"\nAverage choices per respondent: {choices_per_respondent.mean():.2f}")
    
    # Show the legend mapping
    print(f"\n=== LEGEND MAPPING ===")
    legend_mapping = {
        '1': 'Location',
        '2': 'Price', 
        '3': 'Brand Reputation',
        '4': 'Amenities (gym, pool, etc.)',
        '5': 'In-room Entertainment',
        '6': 'Family-friendly Features',
        '7': 'Guest Reviews',
        '99': 'Other'
    }
    
    for code, label in legend_mapping.items():
        print(f"  {code}: {label}")
    
    # Reconstruct proper responses
    print(f"\n=== PROPER RESPONSE RECONSTRUCTION ===")
    
    # For each respondent, collect their choices
    sample_reconstructed = []
    for idx, row in b1a_df.head(10).iterrows():
        choices = []
        for col in b1a_cols:
            value = row[col]
            if pd.notna(value) and str(int(value)) in legend_mapping:
                choices.append(legend_mapping[str(int(value))])
        sample_reconstructed.append(choices)
        print(f"Respondent {idx+1}: {choices}")

if __name__ == "__main__":
    analyze_b1a_responses()
